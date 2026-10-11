"""Turn a case folder (or zip) of bank documents into plain-text builder sources.

Generic: any Markdown, text, Word, PDF, Excel, PowerPoint, CSV, JSON or YAML file
becomes one UTF-8 source file. Files without extractable text (images) are listed,
not guessed. Nothing here knows any case content.
"""
from datetime import date, datetime
from pathlib import Path
import re
import tempfile
import zipfile
from xml.etree import ElementTree

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".yaml", ".yml", ".mmd", ".mermaid",
                 ".puml", ".svg", ".xml", ".html", ".htm", ".drawio", ".log"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FIRST = ("brief.md", "readme.md")


def read_text(path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def docx_text(path):
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    lines = []

    def paragraph(p):
        style = p.find(W + "pPr/" + W + "pStyle")
        name = style.get(W + "val", "") if style is not None else ""
        parts = []
        for node in p.iter():
            if node.tag == W + "t" and node.text:
                parts.append(node.text)
            elif node.tag == W + "tab":
                parts.append("\t")
            elif node.tag in (W + "br", W + "cr"):
                parts.append("\n")
        text = "".join(parts)
        level = re.search(r"(?i)heading\s*(\d)", name) or re.search(r"(?i)ba.l.k\s*(\d)", name)
        if level and text.strip():
            return "#" * min(int(level.group(1)), 6) + " " + text.strip()
        if p.find(W + "pPr/" + W + "numPr") is not None and text.strip():
            return "- " + text.strip()
        return text

    body = root.find(W + "body")
    for block in (body if body is not None else []):
        if block.tag == W + "p":
            lines.append(paragraph(block))
        elif block.tag == W + "tbl":
            for row in block.iter(W + "tr"):
                cells = [" ".join(paragraph(p) for p in cell.iter(W + "p")).strip().replace("|", "/")
                         for cell in row.findall(W + "tc")]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    return "\n".join(lines)


def pptx_text(path):
    lines = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted((n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                        key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
        for index, name in enumerate(slides, 1):
            root = ElementTree.fromstring(archive.read(name))
            lines.append(f"## Slide {index}")
            for p in root.iter(A + "p"):
                text = "".join(t.text or "" for t in p.iter(A + "t")).strip()
                if text:
                    lines.append(text)
    return "\n".join(lines)


def pdf_text(path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, 1):
        pages.append(f"--- page {index} ---\n" + (page.extract_text() or ""))
    return "\n\n".join(pages)


def cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("|", "/").replace("\n", " ").strip()


def xlsx_text(path):
    from openpyxl import load_workbook
    book = load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for sheet in book.worksheets:
        lines.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [cell(v) for v in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                lines.append("| " + " | ".join(values) + " |")
        lines.append("")
    return "\n".join(lines)


CONVERTERS = {".docx": docx_text, ".pdf": pdf_text, ".xlsx": xlsx_text, ".xlsm": xlsx_text, ".pptx": pptx_text}


def unpack(case):
    """A folder is used as is; a zip is extracted to a temporary folder."""
    case = Path(case)
    if case.is_file() and case.suffix.lower() == ".zip":
        target = Path(tempfile.mkdtemp(prefix="case-"))
        with zipfile.ZipFile(case) as archive:
            for member in archive.namelist():
                resolved = (target / member).resolve()
                if not str(resolved).startswith(str(target.resolve())):
                    raise ValueError("Unsafe path in case zip")
            archive.extractall(target)
        return target
    if not case.is_dir():
        raise ValueError("Case input must be a folder or a .zip file")
    return case


def ordered(files, root):
    def key(path):
        name = path.name.lower()
        return (FIRST.index(name) if name in FIRST else len(FIRST), str(path.relative_to(root)).lower())
    return sorted(files, key=key)


def convert_case(case, output, log=print):
    """Write one .md source per readable file into output; return (sources, skipped, structured).

    structured lists JSON/YAML files so a machine-readable service contract can be used directly."""
    root = unpack(case)
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    files = [p for p in root.rglob("*") if p.is_file() and not p.name.startswith((".", "~$"))
             and "__MACOSX" not in p.parts]
    sources, skipped, structured = [], [], []
    for index, path in enumerate(ordered(files, root), 1):
        relative = path.relative_to(root).as_posix()
        suffix = path.suffix.lower()
        try:
            if suffix in TEXT_SUFFIXES:
                text = read_text(path)
                if suffix in {".json", ".yaml", ".yml"}:
                    structured.append(path)
            elif suffix in CONVERTERS:
                text = CONVERTERS[suffix](path)
            else:
                skipped.append((relative, "no text extractor for this file type"))
                continue
        except Exception as exc:  # one unreadable file must not stop the build
            skipped.append((relative, type(exc).__name__ + ": " + str(exc)[:120]))
            continue
        text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not text:
            skipped.append((relative, "no text content"))
            continue
        slug = re.sub(r"[^A-Za-z0-9]+", "-", relative).strip("-")[:80] or "source"
        target = output / f"{index:02d}-{slug}.md"
        target.write_text(f"Source file: {relative}\n\n{text}\n", encoding="utf-8", newline="\n")
        if target.stat().st_size > 2_000_000:
            content = target.read_text(encoding="utf-8")
            target.write_text(content[:1_900_000] + "\n[truncated]\n", encoding="utf-8", newline="\n")
        sources.append(target)
        log(f"source {target.name}: {relative} ({suffix or 'no suffix'}, {len(text)} chars)")
    for relative, reason in skipped:
        log(f"skipped {relative}: {reason}")
    if not sources:
        raise ValueError("The case contains no readable documents")
    return sources, skipped, structured
